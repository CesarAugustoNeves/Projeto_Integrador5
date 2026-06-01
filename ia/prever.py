import pandas as pd
import pickle
import sys
import numpy as np
from pathlib import Path
import openpyxl

# Caminho pro .exe 
if getattr(sys, 'frozen', False):
    caminho_base = Path(sys._MEIPASS)
else:
    caminho_base = Path(__file__).resolve().parent.parent

caminho_modelo = caminho_base / "ia" / "modelo.pkl"

# Carregando a IA
with open(caminho_modelo, 'rb') as arquivo:
    modelo = pickle.load(arquivo)

def prever_excel(arquivo_ia, sheet_name=None, **kwargs):
    abas_para_processar = [sheet_name] if sheet_name is not None else ["Por_Conta", "Por_CCusto", "Por_Evento"]

    for aba in abas_para_processar:
        try:
            df = pd.read_excel(arquivo_ia, sheet_name=aba)
        except Exception:
            continue 
            
        if df.empty:
            continue

        colunas_treino = [
            'adp_debito', 'adp_credito', 
            'sap_debito', 'sap_credito', 
            'adp_saldo', 'sap_saldo', 'diferenca'
        ]
        
        # Garante que todas existam; se faltar alguma, cria com 0
        for col in colunas_treino:
            if col not in df.columns:
                df[col] = 0.0

        # Separa exatamente as 7 colunas esperadas pela IA
        X = df[colunas_treino].fillna(0)

        # Executa a previsão da IA 
        previsoes = modelo.predict(X)

        # Regra de Negocio
        df['Previsao_Final'] = previsoes
        
        for index, row in df.iterrows():
            diff = abs(row['diferenca']) if pd.notna(row['diferenca']) else 0
            
            if diff < 0.01:
                df.at[index, 'Previsao_Final'] = '1 - OK! '
            else:
                adp_tem = (pd.notna(row['adp_saldo']) and row['adp_saldo'] != 0)
                sap_tem = (pd.notna(row['sap_saldo']) and row['sap_saldo'] != 0)
                
                if adp_tem and not sap_tem:
                    df.at[index, 'Previsao_Final'] = '2 - Apenas ADP'
                elif sap_tem and not adp_tem:
                    df.at[index, 'Previsao_Final'] = '3 - Apenas SAP'
                else:
                    df.at[index, 'Previsao_Final'] = '4 - Diferença em ambos'
                    
        previsoes_atualizadas = df['Previsao_Final'].tolist()

        # Gravação
        wb = openpyxl.load_workbook(arquivo_ia)
        if aba in wb.sheetnames:
            ws = wb[aba]
            
            col_idx = None
            for col in range(1, ws.max_column + 1):
                if ws.cell(row=1, column=col).value == 'Previsão do Modelo':
                    col_idx = col
                    break
            
            if col_idx is None:
                col_idx = ws.max_column + 1
                ws.cell(row=1, column=col_idx).value = 'Previsão do Modelo'
            
            for i, pred in enumerate(previsoes_atualizadas, start=2):
                ws.cell(row=i, column=col_idx).value = pred
                
            wb.save(arquivo_ia)
        wb.close()

    print(f"Previsões aplicadas com sucesso para: {abas_para_processar}")